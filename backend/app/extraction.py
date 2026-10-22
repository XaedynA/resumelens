"""Bounded document extraction; uploaded files are never persisted."""
from io import BytesIO
from zipfile import ZipFile
from fastapi import HTTPException
from pypdf import PdfReader
from docx import Document

MAX_TEXT = 100000

def extract_text(filename, raw):
    name = filename.lower()
    if not name.endswith((".pdf", ".docx", ".txt")):
        raise HTTPException(400, "Only PDF, DOCX, and TXT files are supported.")
    try:
        if name.endswith(".pdf"):
            reader = PdfReader(BytesIO(raw))
            if reader.is_encrypted or len(reader.pages) > 50:
                raise HTTPException(422, "Use an unencrypted PDF with at most 50 pages.")
            parts, length = [], 0
            for page in reader.pages:
                contents = page.get_contents()
                if contents is not None and len(contents.get_data()) > 5 * 1024 * 1024:
                    raise HTTPException(422, "PDF page is too complex.")
                text = page.extract_text() or ""
                length += len(text)
                if length > MAX_TEXT:
                    raise HTTPException(422, "Extracted text exceeds 100,000 characters.")
                parts.append(text)
            text = "\n".join(parts)
        elif name.endswith(".docx"):
            with ZipFile(BytesIO(raw)) as archive:
                if len(archive.infolist()) > 1000 or sum(i.file_size for i in archive.infolist()) > 10 * 1024 * 1024:
                    raise HTTPException(422, "DOCX expanded content is too large.")
            doc = Document(BytesIO(raw))
            parts = [p.text for p in doc.paragraphs]
            parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            text = "\n".join(parts)
        else:
            text = raw.decode("utf-8-sig")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(422, "Could not extract text from this file.") from None
    if not text.strip():
        raise HTTPException(422, "No selectable text found. Scanned PDFs are not supported.")
    if len(text) > MAX_TEXT:
        raise HTTPException(422, "Extracted text exceeds 100,000 characters.")
    return text
