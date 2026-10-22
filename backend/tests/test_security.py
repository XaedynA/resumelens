from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import json
import pytest
from fastapi.testclient import TestClient
from docx import Document
from app import main

@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setattr(main, 'DB', tmp_path/'history.sqlite3')
    with TestClient(main.app) as c:
        yield c

@pytest.mark.parametrize('value', ['', ' ', '\n\t', 'x'*100001], ids=['empty', 'spaces', 'whitespace', 'over-limit'])
def test_invalid_input(client, value):
    assert client.post('/api/analyze', json={'resume_text':value,'job_description':'Python'}).status_code == 422

@pytest.mark.parametrize('name,data,status', [('x.exe', b'bad',400), ('x.pdf',b'bad',422), ('x.txt',b'\xff',422), ('x.txt',b' ',422), ('x.txt',b'x'*2097153,413), ('x.txt',b'x'*100001,422)], ids=['unsupported-type', 'invalid-pdf', 'invalid-utf8', 'blank-text', 'file-too-large', 'text-too-long'])
def test_invalid_uploads(client,name,data,status):
    assert client.post('/api/extract',files={'file':(name,data)}).status_code == status

def test_docx_table(client):
    doc=Document();doc.add_table(rows=1, cols=1).cell(0,0).text='Python SQL'
    buf=BytesIO();doc.save(buf)
    assert client.post('/api/extract',files={'file':('r.docx',buf.getvalue())}).json()['text']=='Python SQL'

def test_zip_bomb(client):
    buf=BytesIO()
    with ZipFile(buf,'w',ZIP_DEFLATED) as z:
        z.writestr('word/document.xml',b'x'*(11*1024*1024))
    assert client.post('/api/extract', files={'file':('r.docx',buf.getvalue())}).status_code==422

def test_history_privacy_and_delete(client):
    payload={'resume_text':'Python SQL Docker PRIVATE_MARKER', 'job_description':'Python SQL Docker', 'save':False}
    assert client.post('/api/analyze',json=payload).status_code==200
    assert client.get('/api/history').json()==[]
    payload['save']=True
    assert client.post('/api/analyze',json=payload).status_code==200
    with main.connect() as db:
        result=db.execute('SELECT result_json FROM analyses').fetchone()[0]
    assert 'PRIVATE_MARKER' not in result
    assert 'requirements' not in json.loads(result)
    assert client.delete('/api/history').status_code==204
    assert client.get('/api/history').json()==[]
